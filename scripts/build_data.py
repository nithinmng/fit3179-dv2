"""Build every chart-ready data file for "Outside is Australia's biggest gym" (FIT3179 DV2).

Reads the raw sources in data/raw/ and writes small JSON / GeoJSON files to data/. The web page only
loads the files in data/, never the raw workbooks.

Sources (all cited on the page):
  - AusPlay, Jan to Dec 2025 (Australian Sports Commission, released 30 Apr 2026): data/raw/ausplay/
  - BOM monthly climate statistics for the 8 capitals: data/raw/bom/bom_capitals_monthly.csv
  - OpenStreetMap leisure=fitness_station, snapshot 24 Oct 2026 (ODbL): data/raw/osm/
  - ABS Regional population 2024-25 by LGA 2025, with boundaries: data/raw/abs_lga/
  - ABS ASGS Edition 3 allocation files (mesh block land use 2021, mesh block to LGA 2025): data/raw/abs_mb/
  - ABS SEIFA 2021 by LGA: data/raw/seifa/
  - ABS ASGS 2021 state and territory boundaries (CC BY 4.0): data/raw/states/abs_ste_2021.geojson

Needs Python 3 with openpyxl, matplotlib and numpy (all ship with Anaconda).

    python scripts/build_data.py
"""
from __future__ import annotations

import csv
import json
import math
import statistics
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
from matplotlib.path import Path as MplPath
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
OUT = ROOT / "data"

AGES = ["15-17", "18-24", "25-34", "35-44", "45-54", "55-64", "65+"]
AGE_LABEL = {a: a.replace("-", "–") for a in AGES}
STATES = {
    "NSW": "New South Wales", "VIC": "Victoria", "QLD": "Queensland", "WA": "Western Australia",
    "SA": "South Australia", "TAS": "Tasmania", "ACT": "Australian Capital Territory", "NT": "Northern Territory",
}
# Where each state's spike stands (inland label points, chosen by hand so spikes don't collide)
STATE_ANCHOR = {
    "NSW": (147.2, -32.4), "VIC": (144.3, -36.9), "QLD": (144.8, -22.8), "WA": (121.8, -26.0),
    "SA": (135.6, -29.8), "TAS": (146.7, -42.1), "ACT": (149.9, -35.6), "NT": (133.4, -19.6),
}
CAPITALS = {
    "Greater Sydney": "Sydney", "Greater Melbourne": "Melbourne", "Greater Brisbane": "Brisbane",
    "Greater Perth": "Perth", "Greater Adelaide": "Adelaide", "Greater Hobart": "Hobart",
    "Greater Darwin": "Darwin", "Australian Capital Territory": "Canberra",
}
# Neighbourhood parkland: ABS "Parkland" mesh blocks no larger than this (km2). Larger parkland mesh
# blocks are mostly regional parks, bushland reserves and national parks. See the methods note on the page.
NEIGHBOURHOOD_PARK_MAX_KM2 = 0.5
SITE_MERGE_M = 60          # OSM features closer than this are treated as one outdoor gym site
HEX_RADIUS_KM = 1.5        # hexbin size for the city maps


# Plain-English labels for AusPlay column names (left: start of the AusPlay name)
LABELS = {
    "Physical health": "Physical health or fitness", "Fun": "Fun and enjoyment",
    "To lose weight": "Lose or manage weight", "To be outdoors": "Be outdoors, enjoy nature",
    "Social": "Social reasons", "Psychological": "Mental health", "Sense of achievement": "Sense of achievement",
    "Hobby": "Hobby", "Walk the dog": "Walk the dog", "Training": "Training for an event or sport",
    "Learn a new skill": "Learn a new skill", "Performance": "Performance or competition",
    "Physio": "Rehab or physio", "Professional": "Part of my job", "To be a good role model": "Be a role model",
    "Way of getting around": "Getting around",
    "Apps for music": "Music apps", "Headphones": "Headphones or speakers",
    "Apps for tracking": "Activity-tracking apps", "Wearable": "Wearables (watches, trackers)",
    "Apps for maps": "Maps and GPS apps", "Watching": "Follow-along exercise videos",
    "Indoor smart": "Smart trainers or treadmills", "Apps used during": "Apps used during the activity",
    "Poor health": "Poor health or injury", "Increasing age": "Age, or feeling too old", "Disability": "Disability",
    "Too lazy": "Too lazy", "Not enough time": "Not enough time", "Don": "Don't like sport or exercise",
    "Not a priority": "Not a priority", "Nobody to do": "Nobody to do it with", "Can't afford": "Can't afford it",
    "Fear of injury": "Fear of injury", "Physical job": "Physical job already", "The weather": "The weather",
    "Not good enough": "Not good enough", "No opportunities": "No opportunities or facilities nearby",
    "No transport": "No transport, can't get there", "Not value for money": "Not value for money",
    "1+ per year": "Once a year", "1+ per month": "Monthly", "1+ per fortnight": "Fortnightly",
    "1+ per week": "Weekly", "2+ per week": "2+ a week", "3+ per week": "3+ a week", "4+ per week": "4+ a week",
    "5+ per week": "5+ a week", "6+ per week": "6+ a week", "7+ per week": "7+ a week",
}


def nice(name: str) -> str:
    return next((v for k, v in LABELS.items() if name.startswith(k)), name)


def write(name: str, value) -> None:
    text = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    (OUT / name).write_text(text, encoding="utf-8")
    print(f"  wrote {name:32} {len(text) / 1024:7.1f} KB")


def r4(x: float) -> float:
    return round(x, 4)


# ----------------------------------------------------------------------------------------------
# AusPlay
# ----------------------------------------------------------------------------------------------

def sheet_rows(path: Path, name: str) -> list[list]:
    wb = load_workbook(path, read_only=True, data_only=True)
    return [list(r) for r in wb[name].iter_rows(values_only=True)]


def clean(text) -> str:
    return " ".join(str(text).replace("^", "").replace("#", "").split()).strip()


def rate_block(rows: list[list]) -> list[dict]:
    """Rows of the first rate block in an AusPlay table: [{group, sub, values{column: rate}}].

    The header is the row just above the 'Participation rate' / 'Proportion of ...' marker row; the
    block ends at the 'Estimated number' marker. Column 0 holds the group (All, Males, Females),
    column 1 holds Total or an age group.
    """
    marker = next(i for i, r in enumerate(rows)
                  if any(isinstance(c, str) and (c.startswith("Participation rate") or c.startswith("Proportion"))
                         for c in r[:4]))
    header = rows[marker - 1]
    cols = {j: clean(c) for j, c in enumerate(header) if j >= 2 and c not in (None, "")}
    out, group = [], None
    for r in rows[marker + 1:]:
        if any(isinstance(c, str) and c.startswith("Estimated") for c in r[:4]):
            break
        if r[0] not in (None, "", "Age group"):
            group = clean(r[0])
        if r[1] in (None, ""):
            continue
        vals = {name: r[j] for j, name in cols.items() if isinstance(r[j], (int, float))}
        if vals:
            out.append({"group": group, "sub": clean(r[1]), "values": vals})
    return out


def pick(block, group, sub):
    return next(b["values"] for b in block if b["group"] == group and b["sub"] == sub)


def ausplay():
    nat = RAW / "ausplay" / "national_2025.xlsx"

    # Venues (Table 7)
    t7 = rate_block(sheet_rows(nat, "7"))
    venue_name = {  # column-name prefix -> label on the page
        "Public space": "Public space", "Gym": "Gym or leisure centre", "Home": "Home",
        "Free community": "Free community facility", "Sports club": "Sports club",
    }
    total = pick(t7, "All", "Total")
    venues = []
    for raw_name, rate in total.items():
        label = next((v for k, v in venue_name.items() if raw_name.startswith(k)), None)
        if label:
            venues.append({"venue": label, "rate": r4(rate)})
    venues.sort(key=lambda v: -v["rate"])
    waffle = []
    for order, v in enumerate(venues):
        filled = round(v["rate"] * 100)
        for i in range(100):
            waffle.append({"venue": v["venue"], "order": order, "pct": round(v["rate"] * 100),
                           "i": i, "col": i % 10, "row": i // 10, "filled": i < filled})
    write("venues_waffle.json", waffle)

    pyramid = []
    for sex in ("Males", "Females"):
        for a in AGES:
            vals = pick(t7, sex, a)
            ps = next(v for k, v in vals.items() if k.startswith("Public space"))
            pyramid.append({"sex": "Men" if sex == "Males" else "Women", "age": AGE_LABEL[a], "rate": r4(ps)})
    write("public_space_age_sex.json", pyramid)
    write("overall_age.json", [{"age": AGE_LABEL[a], "rate": r4(pick(t7, "All", a)["Total"])} for a in AGES])

    # Activities by age (Table 17)
    t17_rows = sheet_rows(nat, "17")
    marker = next(i for i, r in enumerate(t17_rows) if r[1] and str(r[1]).startswith("Participation rate"))
    header = t17_rows[marker - 1]
    age_cols = {j: clean(c) for j, c in enumerate(header) if c and clean(c) in AGES}
    acts = []
    for r in t17_rows[marker + 1:]:
        if r[0] is None or not isinstance(r[1], (int, float)):
            if r[0] and str(r[0]).startswith("Estimated"):
                break
            if r[1] and str(r[1]).startswith("Estimated"):
                break
            continue
        acts.append({"activity": clean(r[0]), "total": r[1], "men": r[3], "women": r[4],
                     "ages": {age_cols[j]: r[j] for j in age_cols}})
    acts.sort(key=lambda a: -a["total"])
    short = {"Walking (Recreational)": "Walking", "Fitness/Gym": "Gym and fitness", "Bush walking": "Bushwalking",
             "Running/jogging": "Running", "Football/soccer": "Football (soccer)"}
    act_long = []
    for a in acts[:12]:
        for age, v in a["ages"].items():
            act_long.append({"activity": short.get(a["activity"], a["activity"]), "age": AGE_LABEL[age],
                             "rate": r4(v), "total": r4(a["total"])})
    write("activity_age.json", act_long)

    # Motivations (Table 12)
    t12 = rate_block(sheet_rows(nat, "12"))
    motiv = []
    for b in t12:
        if b["group"] != "All":
            continue
        age = "All ages" if b["sub"] == "Total" else AGE_LABEL.get(b["sub"], b["sub"])
        for name, v in b["values"].items():
            motiv.append({"age": age, "motivation": nice(name), "rate": r4(v)})
    write("motivations.json", motiv)

    # Barriers (Table 13, adults who were not active)
    t13 = rate_block(sheet_rows(nat, "13"))
    barriers = [{"barrier": nice(k), "rate": r4(v)} for k, v in pick(t13, "All", "Total").items()]
    barriers.sort(key=lambda b: -b["rate"])
    keep = barriers[:12] + [b for b in barriers[12:] if b["barrier"] in
                            ("The weather", "No opportunities or facilities nearby", "No transport, can't get there")]
    write("barriers.json", keep)

    # Technology (Table 11)
    t11 = rate_block(sheet_rows(nat, "11"))
    young, old = pick(t11, "All", "18-24"), pick(t11, "All", "65+")
    tech = []
    for name in list(pick(t11, "All", "Total"))[:8]:
        tech.append({"tech": nice(name), "age": "18–24", "rate": r4(young[name])})
        tech.append({"tech": nice(name), "age": "65+", "rate": r4(old[name])})
    write("technology.json", tech)

    # Frequency (Table 9)
    t9 = rate_block(sheet_rows(nat, "9"))
    freq = [{"threshold": nice(k), "rate": r4(v), "order": i} for i, (k, v) in enumerate(pick(t9, "All", "Total").items())]
    write("frequency.json", freq)

    # Income x frequency mosaic (Table 1)
    t1 = sheet_rows(nat, "1")
    start = next(i for i, r in enumerate(t1) if r[0] and str(r[0]).startswith("Annual household income"))
    bands = []
    for r in t1[start:]:
        if r[0] and not str(r[0]).startswith("Annual household income"):
            break
        if r[1] and isinstance(r[2], float) and clean(r[1]) != "Not stated":
            yr, wk, three, count_yr = r[2], r[10], r[18], r[6]
            bands.append({"band": clean(r[1]), "yr": yr, "wk": wk, "three": three, "pop": count_yr / yr})
    total_pop = sum(b["pop"] for b in bands)
    mosaic, x = [], 0.0
    for b in bands:
        w = b["pop"] / total_pop
        classes = [("3+ times a week", b["three"]), ("1–2 times a week", b["wk"] - b["three"]),
                   ("Less than weekly", b["yr"] - b["wk"]), ("Not active in the past year", 1 - b["yr"])]
        y = 0.0
        for k, (name, share) in enumerate(classes):
            mosaic.append({"band": b["band"], "x0": r4(x), "x1": r4(x + w), "y0": r4(y), "y1": r4(y + share),
                           "cls": name, "order": k, "share": r4(share), "weekly": r4(b["wk"]), "width": r4(w)})
            y += share
        x += w
    write("income_mosaic.json", mosaic)

    # States (Table 7 in each state workbook)
    states = []
    for code in STATES:
        rows = sheet_rows(RAW / "ausplay" / f"{code}.xlsx", "7")
        block = rate_block(rows)
        tot = pick(block, "All", "Total")
        hdr_marker = next(i for i, r in enumerate(rows) if any(isinstance(c, str) and c.startswith("Participation rate") for c in r[:4]))
        header = rows[hdr_marker - 1]
        cnt_row = next(r for r in rows if r[0] and str(r[0]).startswith("All#"))
        col = {clean(c): j for j, c in enumerate(header) if c}
        ps_col = next(j for k, j in col.items() if k.startswith("Public space"))
        states.append({
            "state": code, "name": STATES[code],
            "public": r4(next(v for k, v in tot.items() if k.startswith("Public space"))),
            "gym": r4(next(v for k, v in tot.items() if k.startswith("Gym"))),
            "public_people": round(cnt_row[ps_col]),
            "lon": STATE_ANCHOR[code][0], "lat": STATE_ANCHOR[code][1],
        })
    write("states_venues.json", states)
    sankey(nat)
    return states


def sankey(nat: Path):
    """Sankey of age group -> how adults were active (AusPlay Table 3). The four groups do not overlap,
    so every adult is counted once. The ~0.6% who answered "Don't know" are left out (noted on the page).
    Node and link positions are laid out here in pixels for the Vega spec (specs/08-organised-sankey.vg.json)."""
    rows = sheet_rows(nat, "3")
    # Table 3 has two header rows (names, then letter codes), so read columns from the name row
    hdr = rows[next(i for i, r in enumerate(rows) if r[2] == "Total" and r[0] is None)]
    col = {clean(c): j for j, c in enumerate(hdr) if c}
    rate_start = next(i for i, r in enumerate(rows) if r[0] == "All")
    rates = {clean(r[1]): r[col["Total"]] for r in rows[rate_start:rate_start + 8] if r[1] and clean(r[1]) in AGES}
    cnt_start = next(i for i, r in enumerate(rows) if r[0] and str(r[0]).startswith("All#"))
    counts = {}
    for r in rows[cnt_start:cnt_start + 8]:
        if r[1] and clean(r[1]) in AGES:
            counts[clean(r[1])] = r
    cats = [("Informal only", "Informal only"), ("Mix of informal and organised", "Mix of informal"),
            ("Organised only", "Organised only")]
    flows = []
    for a in AGES:
        r = counts[a]
        participants, rate = r[col["Total"]], rates[a]
        population = participants / rate
        for name, prefix in cats:
            j = next(j for k, j in col.items() if k.startswith(prefix))
            flows.append((AGE_LABEL[a], name, r[j], population))
        flows.append((AGE_LABEL[a], "Not active in 2025", population - participants, population))

    W_LEFT, W_RIGHT, NODE_W, TOP, H, GAP_L, GAP_R = 120, 650, 16, 8, 440, 10, 22
    targets = [c[0] for c in cats] + ["Not active in 2025"]
    total = sum(f[2] for f in flows)
    k_left = (H - GAP_L * (len(AGES) - 1)) / total
    k_right = (H - GAP_R * (len(targets) - 1)) / total
    k = min(k_left, k_right)  # one scale for both sides so band widths match node heights
    nodes, y = {}, TOP
    for a in AGES:
        v = sum(f[2] for f in flows if f[0] == AGE_LABEL[a])
        nodes[AGE_LABEL[a]] = {"name": AGE_LABEL[a], "side": "left", "x0": W_LEFT, "x1": W_LEFT + NODE_W,
                               "y0": y, "y1": y + v * k, "value": v}
        y += v * k + GAP_L
    y = TOP
    for t in targets:
        v = sum(f[2] for f in flows if f[1] == t)
        nodes[t] = {"name": t, "side": "right", "x0": W_RIGHT, "x1": W_RIGHT + NODE_W,
                    "y0": y, "y1": y + v * k, "value": v}
        y += v * k + GAP_R
    out_y = {n: nodes[n]["y0"] for n in nodes}
    in_y = dict(out_y)
    links = []
    for a in [AGE_LABEL[x] for x in AGES]:          # source order, then target order within each source
        for t in targets:
            v, pop = next((f[2], f[3]) for f in flows if f[0] == a and f[1] == t)
            w = v * k
            links.append({"source": a, "target": t, "value": round(v), "width": round(w, 2),
                          "share_of_age": r4(v / pop),
                          "sx": W_LEFT + NODE_W, "sy": round(out_y[a] + w / 2, 2),
                          "tx": W_RIGHT, "ty": 0.0})
            out_y[a] += w
    for t in targets:                                 # stack each target's incoming bands in source order
        for l in links:
            if l["target"] == t:
                l["ty"] = round(in_y[t] + l["width"] / 2, 2)
                in_y[t] += l["width"]
    adults = sum(n["value"] for n in nodes.values() if n["side"] == "left")
    for n in nodes.values():
        n["share"] = r4(n["value"] / adults)
        n["label"] = n["name"]
        n["sublabel"] = (f"{n['value'] / 1e6:.1f}M adults" if n["side"] == "left"
                         else f"{n['value'] / 1e6:.1f}M · {round(100 * n['value'] / adults)}% of adults")
        for key in ("y0", "y1"):
            n[key] = round(n[key], 2)
        n["value"] = round(n["value"])
    write("sankey_organised.json", {"nodes": list(nodes.values()), "links": links})


# ----------------------------------------------------------------------------------------------
# BOM
# ----------------------------------------------------------------------------------------------

def bom():
    use = {"Sydney Airport": "Sydney"}  # current station (1939 to 2026); Observatory Hill ends in 2020
    skip = {"Sydney"}
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    days = [31, 28.25, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    data = defaultdict(dict)
    meta = {}
    with open(RAW / "bom" / "bom_capitals_monthly.csv", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if row["city"] in skip:
                continue
            city = use.get(row["city"], row["city"])
            meta.setdefault(city, {"station": f"{row['station_name']} ({row['station_id']})"})[row["element"]] = row["period"]
            for i, m in enumerate(months):
                data[(city, i)][row["element"]] = float(row[m])
    out = []
    for (city, i), e in sorted(data.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        for kind, key in (("Hot days (30 °C or more)", "days_ge_30c"), ("Rain days (1 mm or more)", "rain_days_ge_1mm"),
                          ("Cold mornings (2 °C or less)", "days_le_2c")):
            out.append({"city": city, "month": months[i], "m": i, "kind": kind, "days": e[key],
                        "share": r4(e[key] / days[i]), "max_c": e["mean_max_c"],
                        "station": meta[city]["station"], "years_temp": meta[city]["mean_max_c"],
                        "years_rain": meta[city]["rain_days_ge_1mm"]})
    write("climate_rose.json", out)


# ----------------------------------------------------------------------------------------------
# Geometry helpers (no GIS libraries needed)
# ----------------------------------------------------------------------------------------------

def ring_area(ring):
    return sum(x0 * y1 - x1 * y0 for (x0, y0), (x1, y1) in zip(ring, ring[1:])) / 2


def simplify(ring, tol):
    """Douglas-Peucker on one closed ring (lon/lat degrees). The ring is split at the vertex farthest
    from its start, so each half is an open path with distinct end points."""
    if len(ring) < 5:
        return ring
    far = max(range(len(ring)), key=lambda i: (ring[i][0] - ring[0][0]) ** 2 + (ring[i][1] - ring[0][1]) ** 2)
    keep = [False] * len(ring)
    keep[0] = keep[far] = keep[-1] = True
    stack = [(0, far), (far, len(ring) - 1)]
    while stack:
        a, b = stack.pop()
        (x0, y0), (x1, y1) = ring[a], ring[b]
        dx, dy = x1 - x0, y1 - y0
        norm = math.hypot(dx, dy) or 1e-12
        best, idx = 0.0, None
        for i in range(a + 1, b):
            d = abs(dy * ring[i][0] - dx * ring[i][1] + x1 * y0 - y1 * x0) / norm
            if d > best:
                best, idx = d, i
        if idx is not None and best > tol:
            keep[idx] = True
            stack += [(a, idx), (idx, b)]
    return [p for p, k in zip(ring, keep) if k]


def polygons_of(geom):
    return [geom["coordinates"]] if geom["type"] == "Polygon" else geom["coordinates"]


def clean_geometry(geom, tol, min_area, decimals=4):
    """Simplify, drop tiny islands, round to 4 decimals and wind rings the way d3/Vega expects
    (exterior rings clockwise, holes anticlockwise)."""
    polys = []
    for poly in polygons_of(geom):
        rings = []
        for k, ring in enumerate(poly):
            s = simplify(ring, tol)
            if len(s) < 4 or abs(ring_area(s)) < (min_area if k == 0 else min_area * 4):
                if k == 0:
                    break
                continue
            s = [[round(x, decimals), round(y, decimals)] for x, y in s]
            area = ring_area(s)
            if (k == 0 and area > 0) or (k > 0 and area < 0):
                s.reverse()
            rings.append(s)
        if rings:
            polys.append(rings)
    if not polys:
        return None
    return {"type": "MultiPolygon", "coordinates": polys}


def centroid(geom):
    best = max(polygons_of(geom), key=lambda p: abs(ring_area(p[0])))
    ring = best[0]
    a = ring_area(ring)
    cx = sum((x0 + x1) * (x0 * y1 - x1 * y0) for (x0, y0), (x1, y1) in zip(ring, ring[1:])) / (6 * a)
    cy = sum((y0 + y1) * (x0 * y1 - x1 * y0) for (x0, y0), (x1, y1) in zip(ring, ring[1:])) / (6 * a)
    return round(cx, 4), round(cy, 4)


def state_features():
    """ABS 2021 state and territory boundaries, without 'Other Territories' and 'Outside Australia'."""
    d = json.loads((RAW / "states" / "abs_ste_2021.geojson").read_text(encoding="utf-8"))
    return [f for f in d["features"] if f.get("geometry") and f["properties"]["state_name_2021"] in STATES.values()]


def states_geo():
    abbr = {v: k for k, v in STATES.items()}
    feats = []
    for f in state_features():
        name = f["properties"]["state_name_2021"]
        g = clean_geometry(f["geometry"], 0.02, 0.05, decimals=3)
        feats.append({"type": "Feature", "properties": {"state": abbr[name], "name": name}, "geometry": g})
    write("states.json", {"type": "FeatureCollection", "features": feats})


# ----------------------------------------------------------------------------------------------
# Councils: population, neighbourhood parkland, SEIFA
# ----------------------------------------------------------------------------------------------

def councils():
    print("  reading mesh block allocation files (takes a minute)...")
    ws = load_workbook(RAW / "abs_mb" / "MB_2021_AUST.xlsx", read_only=True, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    ix = {h: i for i, h in enumerate(next(it))}
    mb = {}
    for r in it:
        if r[0] is not None:
            mb[str(r[ix["MB_CODE_2021"]])] = (r[ix["MB_CATEGORY_2021"]], r[ix["AREA_ALBERS_SQKM"]] or 0.0,
                                               r[ix["GCCSA_NAME_2021"]] or "")
    ws = load_workbook(RAW / "abs_mb" / "LGA_2025_AUST.xlsx", read_only=True, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    ix = {h: i for i, h in enumerate(next(it))}
    park_local, park_all, area_all = Counter(), Counter(), Counter()
    city_area = defaultdict(Counter)
    for r in it:
        if r[0] is None:
            continue
        cat, area, gcc = mb.get(str(r[ix["MB_CODE_2021"]]), (None, 0.0, ""))
        lga = str(r[ix["LGA_CODE_2025"]])
        area_all[lga] += area
        if gcc in CAPITALS:
            city_area[lga][CAPITALS[gcc]] += area
        if cat == "Parkland":
            park_all[lga] += area
            if area <= NEIGHBOURHOOD_PARK_MAX_KM2:
                park_local[lga] += area

    seifa_ws = load_workbook(RAW / "seifa" / "seifa_2021_lga.xlsx", read_only=True, data_only=True)["Table 1"]
    seifa = {str(r[0]): (r[2], r[3]) for r in seifa_ws.iter_rows(values_only=True)
             if r and r[0] is not None and str(r[0]).isdigit()}
    seifa["24700"] = seifa.get("25250")  # Merri-bek (2025) was Moreland (2021 code 25250)

    lga = json.loads((RAW / "abs_lga" / "lga_erp_2025.geojson").read_text(encoding="utf-8"))
    abbr = {v: k for k, v in STATES.items()}
    rows, city_feats = [], []
    for f in lga["features"]:
        p = f["properties"]
        code, pop = p["lga_code_2025"], p["erp_2025"] or 0
        if not pop or not f.get("geometry"):
            continue
        city = None
        if city_area[code]:
            name, a = city_area[code].most_common(1)[0]
            if a / max(area_all[code], 1e-9) >= 0.5:
                city = name
        s = seifa.get(code)
        cx, cy = centroid(f["geometry"])
        row = {
            "code": code, "council": p["lga_name_2025"], "state": abbr.get(p["state_name_2021"], "OT"),
            "city": city, "pop": pop, "density": round(p["pop_density_2025_people_per_km2"] or 0, 1),
            "park_local_km2": round(park_local[code], 3), "park_all_km2": round(park_all[code], 2),
            "park_m2_pp": round(1e6 * park_local[code] / pop, 1),
            "park_all_m2_pp": round(1e6 * park_all[code] / pop),
            "park_share": round(100 * park_all[code] / max(area_all[code], 1e-9), 1),
            "irsd": s[0] if s else None, "irsd_decile": s[1] if s else None,
            "lon": cx, "lat": cy,
        }
        rows.append(row)
        if city:
            g = clean_geometry(f["geometry"], 0.0012, 1e-6, decimals=3)
            if g:
                keep = {k: row[k] for k in ("council", "city", "pop", "density", "park_m2_pp", "park_all_m2_pp")}
                city_feats.append({"type": "Feature", "properties": keep, "geometry": g})

    # Density bands within the capital-city councils (thirds), used for colour in the beeswarm
    cap = [r for r in rows if r["city"]]
    cuts = statistics.quantiles([r["density"] for r in cap], n=3)
    for r in rows:
        r["density_band"] = ("Dense (inner)" if r["density"] >= cuts[1] else
                             "Middle" if r["density"] >= cuts[0] else "Spread out (outer)")
        d = r["irsd_decile"]
        r["irsd_quintile"] = None if d is None else (d + 1) // 2
    write("councils.json", rows)
    write("city_councils.json", {"type": "FeatureCollection", "features": city_feats})

    # Beeswarm layout for capital-city councils: x is log10(m2 per resident) over a fixed pixel width,
    # y stacks dots within each SEIFA quintile row without overlaps.
    W, R, ROW_H = 820, 5.2, 92
    lo, hi = math.log10(3), math.log10(600)
    swarm = []
    for q in range(1, 6):
        pts = sorted((r for r in cap if r["irsd_quintile"] == q and r["park_m2_pp"] > 0), key=lambda r: r["park_m2_pp"])
        placed = []
        for r in pts:
            x = (math.log10(max(min(r["park_m2_pp"], 600), 3)) - lo) / (hi - lo) * W
            dy = 0.0
            for step in range(0, 200):
                dy = (step + 1) // 2 * (1 if step % 2 else -1) * R * 0.9
                if all((x - px) ** 2 + (dy - py) ** 2 >= (2 * R) ** 2 for px, py in placed):
                    break
            placed.append((x, dy))
            swarm.append({**{k: r[k] for k in ("council", "city", "pop", "density", "density_band", "park_m2_pp",
                                                 "irsd_decile", "irsd_quintile")},
                          "y": round((q - 1) * ROW_H + ROW_H / 2 + dy, 1)})
    write("council_swarm.json", swarm)

    # Sensitivity of the headline finding to the size cut-off is reported on the page; print the
    # rank correlations for the chosen cut-off here as a check.
    def ranks(v):  # average ranks for ties
        order = sorted(range(len(v)), key=lambda i: v[i])
        out, i = [0.0] * len(v), 0
        while i < len(v):
            j = i
            while j + 1 < len(v) and v[order[j + 1]] == v[order[i]]:
                j += 1
            for k in range(i, j + 1):
                out[order[k]] = (i + j) / 2
            i = j + 1
        return out

    def spearman(a, b):
        xa, xb = ranks(a), ranks(b)
        ma, mb_ = statistics.mean(xa), statistics.mean(xb)
        num = sum((p - ma) * (q - mb_) for p, q in zip(xa, xb))
        return num / math.sqrt(sum((p - ma) ** 2 for p in xa) * sum((q - mb_) ** 2 for q in xb))
    s_rows = [r for r in cap if r["irsd"]]
    print(f"  capital-city councils: {len(cap)}; Spearman park vs IRSD "
          f"{spearman([r['park_m2_pp'] for r in s_rows], [r['irsd'] for r in s_rows]):.2f}; "
          f"park vs density {spearman([r['park_m2_pp'] for r in cap], [r['density'] for r in cap]):.2f}")
    return rows, city_feats, lga


# ----------------------------------------------------------------------------------------------
# OpenStreetMap outdoor gyms: merge into sites, assign to councils, hexbins per capital
# ----------------------------------------------------------------------------------------------

def dist_m(a, b):
    dy = (a[1] - b[1]) * 111_320
    dx = (a[0] - b[0]) * 111_320 * math.cos(math.radians(a[1]))
    return math.hypot(dx, dy)


def osm(rows, lga):
    raw = json.loads((RAW / "osm" / "fitness_stations_raw.json").read_text(encoding="utf-8"))
    feats = []
    for el in raw["elements"]:
        lat = el.get("lat", el.get("center", {}).get("lat"))
        lon = el.get("lon", el.get("center", {}).get("lon"))
        if lat is not None:
            feats.append((lon, lat, el.get("tags", {}).get("name", "")))
    feats.sort(key=lambda f: (f[1], f[0]))
    grid, sites = defaultdict(list), []
    for lon, lat, name in feats:
        key = (round(lat * 500), round(lon * 500))
        near = [s for k in [(key[0] + i, key[1] + j) for i in (-1, 0, 1) for j in (-1, 0, 1)] for s in grid[k]]
        hit = next((s for s in near if dist_m(sites[s][:2], (lon, lat)) <= SITE_MERGE_M), None)
        if hit is None:
            sites.append([lon, lat, name])
            grid[key].append(len(sites) - 1)
        elif name and not sites[hit][2]:
            sites[hit][2] = name

    by_code = {r["code"]: r for r in rows}
    polys = []
    for f in lga["features"]:
        code = f["properties"]["lga_code_2025"]
        if code not in by_code or not f.get("geometry"):
            continue
        for poly in polygons_of(f["geometry"]):
            ring = poly[0]
            xs, ys = [c[0] for c in ring], [c[1] for c in ring]
            polys.append((min(xs), min(ys), max(xs), max(ys), MplPath(ring), by_code[code]))
    out = []
    for lon, lat, name in sites:
        r = next((p for x0, y0, x1, y1, path, p in polys
                  if x0 <= lon <= x1 and y0 <= lat <= y1 and path.contains_point((lon, lat))), None)
        site = {"lon": round(lon, 4), "lat": round(lat, 4), "council": r["council"] if r else "",
                "state": r["state"] if r else "", "city": (r or {}).get("city")}
        if name:
            site["name"] = name
        out.append(site)
    # The page needs position, name, council and state; city is only used here for the hexbins
    write("osm_sites.json", {"snapshot": raw["osm_timestamp"],
                             "sites": [{k: v for k, v in s.items() if k != "city"} for s in out]})

    # Hexbins (pointy-top) per capital in a local equirectangular plane
    hexes = []
    size = HEX_RADIUS_KM
    for city in sorted({s["city"] for s in out if s["city"]}):
        pts = [s for s in out if s["city"] == city]
        lat0 = statistics.mean(s["lat"] for s in pts)
        kx, ky = 111.32 * math.cos(math.radians(lat0)), 110.574
        counts = Counter()
        for s in pts:
            x, y = s["lon"] * kx, s["lat"] * ky
            q = (math.sqrt(3) / 3 * x - y / 3) / size
            r_ = (2 / 3 * y) / size
            rx, ry, rz = round(q), round(r_), round(-q - r_)
            dx, dy, dz = abs(rx - q), abs(ry - r_), abs(rz + q + r_)
            if dx > dy and dx > dz:
                rx = -ry - rz
            elif dy > dz:
                ry = -rx - rz
            counts[(rx, ry)] += 1
        for (q, r_), n in counts.items():
            cx = size * math.sqrt(3) * (q + r_ / 2)
            cy = size * 1.5 * r_
            ring = []
            for k in range(6, -1, -1):  # clockwise for d3
                ang = math.radians(60 * k - 30)
                ring.append([round((cx + size * math.cos(ang)) / kx, 5), round((cy + size * math.sin(ang)) / ky, 5)])
            hexes.append({"type": "Feature", "properties": {"city": city, "sites": n},
                          "geometry": {"type": "Polygon", "coordinates": [ring]}})
    write("osm_city_hexes.json", {"type": "FeatureCollection", "features": hexes})
    print(f"  OSM sites: {len(out)}; in capital-city councils: {sum(1 for s in out if s['city'])}")


# ----------------------------------------------------------------------------------------------
# Dorling cartogram layout (chart 17)
# ----------------------------------------------------------------------------------------------

def conic_equal_area(parallels=(-18.0, -36.0), rotate=-132.0):
    """d3's conicEqualArea (Albers) with the same parallels and rotation as the Vega-Lite maps.
    Returns (x, y) with y pointing down, before scaling."""
    p0, p1 = math.radians(parallels[0]), math.radians(parallels[1])
    sy0 = math.sin(p0)
    n = (sy0 + math.sin(p1)) / 2
    c = 1 + sy0 * (2 * n - sy0)
    r0 = math.sqrt(c) / n

    def project(lon, lat):
        lam = math.radians(((lon + rotate + 180) % 360) - 180)
        r = math.sqrt(c - 2 * n * math.sin(math.radians(lat))) / n
        return r * math.sin(lam * n), -(r0 - r * math.cos(lam * n))
    return project


def dorling(rows):
    """Circles sized by population (area), placed at each council's projected centroid, then pushed
    apart until they no longer overlap (bigger circles move less), with a gentle pull back home.
    Computed here so the page does not run a force simulation while it loads."""
    W, H, PAD, BOTTOM = 860, 600, 12, 100
    project = conic_equal_area()
    states = state_features()
    pts = [project(x, y) for f in states for poly in polygons_of(f["geometry"]) for x, y in poly[0]]
    x0, x1 = min(p[0] for p in pts), max(p[0] for p in pts)
    y0, y1 = min(p[1] for p in pts), max(p[1] for p in pts)
    k = min((W - 2 * PAD) / (x1 - x0), (H - BOTTOM - 2 * PAD) / (y1 - y0))
    ox = (W - k * (x1 - x0)) / 2 - k * x0
    oy = PAD + (H - BOTTOM - 2 * PAD - k * (y1 - y0)) / 2 - k * y0

    def to_px(lon, lat):
        x, y = project(lon, lat)
        return ox + k * x, oy + k * y

    abbr = {v: s for s, v in STATES.items()}
    shapes = []
    for f in states:
        g = clean_geometry(f["geometry"], 0.02, 0.05)
        d = ""
        for poly in g["coordinates"]:
            for ring in poly:
                d += "M" + "L".join(f"{x:.1f},{y:.1f}" for x, y in (to_px(*v) for v in ring)) + "Z"
        shapes.append({"state": abbr[f["properties"]["state_name_2021"]], "path": d})

    people = [r for r in rows if r["pop"] > 0]
    total = sum(r["pop"] for r in people)
    s = math.sqrt(0.15 * W * (H - BOTTOM) / (math.pi * total))   # circles cover about 15% of the map
    home = np.array([to_px(r["lon"], r["lat"]) for r in people])
    rad = np.array([math.sqrt(r["pop"]) * s for r in people])
    pos = home.copy()
    for step in range(900):  # the last 300 steps only separate circles, with no pull home
        d = pos[:, None, :] - pos[None, :, :]
        dist = np.sqrt((d ** 2).sum(-1)) + 1e-9
        overlap = np.clip(rad[:, None] + rad[None, :] + 0.8 - dist, 0, None)
        np.fill_diagonal(overlap, 0)
        share = rad[None, :] ** 2 / (rad[:, None] ** 2 + rad[None, :] ** 2)
        pos += ((d / dist[..., None]) * (overlap * share)[..., None] * 0.5).sum(1)
        if step < 600:
            pos += (home - pos) * 0.03
    d = pos[:, None, :] - pos[None, :, :]
    gap = np.sqrt((d ** 2).sum(-1)) - rad[:, None] - rad[None, :]
    np.fill_diagonal(gap, 1e9)
    worst = float(-min(gap.min(), 0))

    circles = []
    for r, (x, y), rr in zip(people, pos, rad):
        circles.append({"council": "Canberra (ACT)" if r["council"] == "Unincorporated ACT" else r["council"],
                        "state": r["state"], "city": r["city"], "pop": r["pop"], "park": r["park_m2_pp"],
                        "park_all": r["park_all_m2_pp"], "zero": r["park_m2_pp"] == 0,
                        "x": round(float(x), 1), "y": round(float(y), 1), "r": round(float(rr), 2)})

    # Capital-city names at each cluster's population-weighted centre (drawn with a halo), then council
    # labels for the biggest circles, skipping any that would collide with a label already placed
    cities, boxes = [], []
    for city in CAPITALS.values():
        cl = [c_ for c_ in circles if c_["city"] == city]
        if cl:
            wsum = sum(c_["pop"] for c_ in cl)
            cx = sum(c_["x"] * c_["pop"] for c_ in cl) / wsum
            cy = sum(c_["y"] * c_["pop"] for c_ in cl) / wsum
            cities.append({"text": city.upper(), "x": round(cx, 1), "y": round(cy, 1)})
            w = len(city) * 8.4
            boxes.append((cx - w / 2, cy - 8, cx + w / 2, cy + 8))
    labels = []
    for c_ in sorted(circles, key=lambda c: -c["pop"]):
        if c_["pop"] < 300_000:
            break
        w, h = len(c_["council"]) * 6.3, 13
        box = (c_["x"] - w / 2, c_["y"] - h / 2, c_["x"] + w / 2, c_["y"] + h / 2)
        if all(box[2] < b[0] or box[0] > b[2] or box[3] < b[1] or box[1] > b[3] for b in boxes):
            boxes.append(box)
            labels.append({"text": c_["council"], "x": c_["x"], "y": c_["y"], "dark": c_["park"] >= 120})

    legend = [{"label": lab, "pop": p, "r": round(math.sqrt(p) * s, 2)}
              for lab, p in (("100,000", 100_000), ("500,000", 500_000), ("1 million", 1_000_000))]
    write("dorling.json", {"width": W, "height": H, "states": shapes, "circles": circles, "labels": labels,
                           "cities": cities, "size_legend": legend})
    print(f"  Dorling: {len(circles)} circles, worst remaining overlap {worst:.2f}px, {len(labels)} labels")


def main():
    print("AusPlay")
    ausplay()
    print("BOM")
    bom()
    print("Geometry")
    states_geo()
    print("Councils")
    rows, _, lga = councils()
    dorling(rows)
    print("OpenStreetMap")
    osm(rows, lga)


if __name__ == "__main__":
    main()
